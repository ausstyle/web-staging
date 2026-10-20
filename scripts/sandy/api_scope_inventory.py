#!/usr/bin/env python3
"""Inventory repository evidence without claiming API or production readiness.

This is a lexical source inventory and optional original-scope source review,
not a substitute for reading public.sandy_operation_registry or live testing.
Only files below supabase/functions and supabase/migrations are inspected.
"""
import argparse
import csv
import hashlib
import json
import re
import subprocess
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SOURCE_ONLY = "source_only_not_deployed_or_functionally_verified"
LIMITATIONS = [
    "The original 161-item Sandy API Scope Document has not been supplied to this inventory tool; no original requirements are inferred or reconciled.",
    "public.sandy_operation_registry is the authoritative operation catalogue. Repository candidates do not establish its current contents.",
    "Literal action comparisons are lexical source candidates, not proof of executable operations, authorization, reachability or completeness.",
    "Source hashes describe files at inventory time, including working-tree changes. They do not certify production parity or a deployed release.",
    "No live registry, Supabase, Google Cloud, GitHub API or external provider request is made by this tool. Deployment and functional verification remain unverified.",
    "No credential files, process environment values or application data are inspected. XLSX output requires the separately installed openpyxl package.",
]


def _git(root, *args):
    result = subprocess.run(["git", "-C", str(root), *args], capture_output=True,
                            text=True, check=False)
    return result.stdout.strip() if result.returncode == 0 else None


def _without_comments(source):
    """Mask comments while preserving line numbers and quoted source strings."""
    token = re.compile(r'("(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\'|`(?:\\.|[^`\\])*`)|(//[^\n]*|/\*[\s\S]*?\*/)')
    return token.sub(lambda m: m.group(1) if m.group(1) else
                     re.sub(r"[^\n]", " ", m.group(2)), source)


def _literal(value):
    # Restrict names to identifier-like values; do not export arbitrary literals.
    return bool(re.fullmatch(r"[A-Za-z][A-Za-z0-9_.:-]{0,140}", value))


def action_candidates(source):
    """Find action equality checks only; inequality and arbitrary strings do not count."""
    clean = _without_comments(source)
    operand = r"(?<![\w.])(?:action|(?:input|body)\.action)(?![\w.])"
    literal = r"(?P<quote>['\"])(?P<value>[A-Za-z][A-Za-z0-9_.:-]{0,140})(?P=quote)"
    patterns = [re.compile(operand + r"\s*==={0,1}\s*" + literal),
                re.compile(literal + r"\s*==={0,1}\s*" + operand)]
    rows = []
    seen = set()
    for pattern in patterns:
        for match in pattern.finditer(clean):
            value = match.group("value")
            line = clean.count("\n", 0, match.start()) + 1
            if (value, line) not in seen:
                rows.append({"action_candidate": value, "line": line})
                seen.add((value, line))
    return sorted(rows, key=lambda row: (row["line"], row["action_candidate"]))


def gateway_handlers(source):
    """Read literal entries in the explicit allowedHandlers source map only."""
    clean = _without_comments(source)
    start = re.search(r"\bconst\s+allowedHandlers\b[^=]*=\s*(?:Object\.freeze\(\s*)?\{", clean)
    if not start:
        return []
    # The map's next top-level '};' ends this intentionally narrow parser.
    closing = re.search(r"\n\}\)?;", clean[start.end():])
    if closing is None:
        return []
    end = start.end() + closing.start()
    block = clean[start.end():end]
    entry = re.compile(r"['\"](?P<operation>[A-Za-z][A-Za-z0-9_.-]*)['\"]\s*:\s*\{"
                       r"\s*functionName\s*:\s*['\"](?P<function>[a-z][a-z0-9-]*)['\"]\s*,"
                       r"\s*action\s*:\s*['\"](?P<action>[A-Za-z][A-Za-z0-9_.:-]*)['\"]\s*,?\s*\}")
    return [{"operation": match.group("operation"), "handler_function": match.group("function"),
             "handler_action": match.group("action"),
             "line": clean.count("\n", 0, start.end() + match.start()) + 1,
             "evidence_status": SOURCE_ONLY} for match in entry.finditer(block)]


def inventory(root=ROOT):
    root = Path(root).resolve()
    functions_dir = root / "supabase/functions"
    if not functions_dir.is_dir():
        raise ValueError("Repository has no supabase/functions directory")
    files = sorted(path for path in functions_dir.rglob("*")
                   if path.is_file() and path.suffix in {".ts", ".js", ".json"})
    migrations = sorted((root / "supabase/migrations").glob("*.sql"))
    entries, actions, handlers, references, source_files = [], [], [], [], []
    function_names = {path.parent.name for path in files if path.name == "index.ts"
                      and path.parent.parent == functions_dir}
    for path in files + migrations:
        if not path.resolve().is_relative_to(root):
            raise ValueError("Refusing source symlink outside repository")
        data = path.read_bytes()
        relative = path.relative_to(root).as_posix()
        digest = hashlib.sha256(data).hexdigest()
        source_files.append({"path": relative, "sha256": digest, "bytes": len(data)})
        if path in migrations:
            continue
        name = path.relative_to(functions_dir).parts[0]
        if path.name == "index.ts" and path.parent.parent == functions_dir:
            entries.append({"function_slug": name, "path": relative, "sha256": digest,
                            "evidence_status": SOURCE_ONLY})
        if path.suffix not in {".ts", ".js"}:
            continue
        if re.search(r"(?:_test|\.test)\.(?:ts|js)$", path.name) or any(
                part in {"test", "tests", "node_modules"} for part in path.relative_to(functions_dir).parts):
            # Retain test-file hashes for provenance; test fixtures are not API handlers.
            continue
        source = data.decode("utf-8", errors="replace")
        for row in action_candidates(source):
            actions.append({"function_slug": name, "path": relative, **row,
                            "evidence_status": "lexical_candidate_not_an_operation"})
        for row in gateway_handlers(source):
            handlers.append({"path": relative, **row,
                             "handler_source_present": row["handler_function"] in function_names})
        for match in re.finditer(r"/functions/v1/([a-z][a-z0-9-]+)", _without_comments(source)):
            target = match.group(1)
            references.append({"caller_function": name, "target_function": target,
                               "path": relative, "line": source.count("\n", 0, match.start()) + 1,
                               "target_source_present": target in function_names,
                               "evidence_status": "literal_reference_not_execution_proof"})
    missing_workers = [row for row in references if not row["target_source_present"]
                       and re.search(r"worker|provision|executor", row["target_function"])]
    head = _git(root, "rev-parse", "--verify", "HEAD^{commit}")
    state = _git(root, "status", "--short", "--untracked-files=normal")
    return {"schema_version": 1, "inventory_scope": "repository_only",
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "source_root": str(root), "source_head": head,
            "working_tree_status": state.splitlines() if state else [],
            "scope_document_supplied": False, "original_requirement_rows": [],
            "original_requirement_count": None, "requirements_reconciled": False,
            "production_verified": False, "limitations": LIMITATIONS,
            "counts": {"edge_functions": len(entries), "action_candidates": len(actions),
                       "gateway_handlers": len(handlers), "function_references": len(references),
                       "missing_worker_source_references": len(missing_workers),
                       "migrations": len(migrations)},
            "edge_functions": entries, "action_candidates": actions,
            "gateway_handlers": handlers, "function_references": references,
            "missing_worker_sources": missing_workers,
            "migrations": [row for row in source_files if row["path"].startswith("supabase/migrations/")],
            "source_files": source_files}


def _requirement_id(value):
    if type(value) is int and value > 0:
        return value
    if isinstance(value, str):
        match = re.fullmatch(r"(?:(?:REQ|R|API)[-_ ]?)?0*([1-9][0-9]*)", value, re.I)
        if match:
            return int(match.group(1))
    raise ValueError("Invalid original requirement ID: " + str(value))


def _read_markdown_scope(path):
    """Parse explicit consecutive numbered requirements beneath level-two headings."""
    rows = []
    seen = set()
    heading = None
    fence = None
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        fence_match = re.match(r"^\s*(`{3,}|~{3,})", line)
        if fence_match:
            marker = fence_match.group(1)[0]
            if fence is None:
                fence = marker
            elif fence == marker:
                fence = None
            continue
        if fence is not None:
            continue
        section = re.match(r"^##[ \t]+(.+?)\s*$", line)
        if section:
            heading = section.group(1)
            continue
        match = re.match(r"^([1-9][0-9]*)\.[ \t]+(.+)$", line)
        if not match:
            continue
        identity = int(match.group(1))
        if identity in seen:
            raise ValueError("Duplicate Markdown requirement ID: " + str(identity))
        if identity != len(rows) + 1:
            raise ValueError(f"Expected Markdown requirement ID {len(rows) + 1}, found {identity}")
        if heading is None:
            raise ValueError("Markdown requirement has no preceding level-two domain heading")
        seen.add(identity)
        operation, _, description = match.group(2).partition(" — ")
        if not operation.strip():
            raise ValueError("Markdown requirement has no operation text")
        rows.append({"requirement_id": identity, "workbook_row": None,
                     "source_line": line_number, "source_heading": heading,
                     "original_requirement_text": line, "domain": heading,
                     "proposed_api_operation_or_function_family": operation,
                     "required_capability": description, "classification": ""})
    return rows


def _read_workbook_scope(scope_path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading the supplied scope workbook requires openpyxl") from exc
    workbook = load_workbook(scope_path, read_only=True, data_only=False)
    try:
        sheet = workbook["API Function Catalogue"]
        iterator = iter(sheet.values)
        header = next(iterator)
        expected_header = ("Domain", "Proposed API operation / function family", "Required capability", "Classification")
        if tuple(header) != expected_header:
            raise ValueError("Unexpected original scope workbook header")
        rows = []
        for source_row, values in enumerate(iterator, start=2):
            if all(value is None for value in values):
                continue
            if len(values) != 4 or any(not isinstance(value, str) or not value.strip() for value in values):
                raise ValueError("Malformed original requirement at workbook row " + str(source_row))
            rows.append({"requirement_id": len(rows) + 1, "workbook_row": source_row,
                         "source_line": None, "source_heading": None,
                         "original_requirement_text": None,
                         "domain": values[0], "proposed_api_operation_or_function_family": values[1],
                         "required_capability": values[2], "classification": values[3]})
    finally:
        workbook.close()
    return rows


def attach_scope(report, scope_path, reconciliation_dir=None, expected_count=161):
    """Preserve original Markdown or workbook rows and merge source review only.

    IDs are explicit Markdown numbers or one-based data-row order in the named
    original workbook sheet. JSON evidence cannot replace original text or mark
    deployments or tests verified. Markdown reading uses only the stdlib.
    """
    scope_path = Path(scope_path).resolve()
    digest = hashlib.sha256(scope_path.read_bytes()).hexdigest()
    suffix = scope_path.suffix.lower()
    if suffix == ".md":
        rows = _read_markdown_scope(scope_path)
        document = {"path": str(scope_path), "format": "markdown", "sha256": digest}
    elif suffix == ".xlsx":
        rows = _read_workbook_scope(scope_path)
        document = {"path": str(scope_path), "format": "xlsx", "sheet": "API Function Catalogue", "sha256": digest}
    else:
        raise ValueError("Scope must be a canonical .md file or original .xlsx workbook")
    if len(rows) != expected_count:
        raise ValueError(f"Expected {expected_count} original requirements, found {len(rows)}")
    for row in rows:
        row.update({"scope_document_sha256": digest, "source_status": "not_reviewed",
                    "source_paths": [], "source_evidence": [], "missing_capability": "Not assessed",
                    "required_access": [], "live_registry_mapping": "unverified",
                    "production_deployment": "unverified", "functional_verification": "unverified",
                    "sandy_mcp_callability": "unverified", "production_verified": False})
    seen = set()
    part_sources = []
    if reconciliation_dir is not None:
        directory = Path(reconciliation_dir)
        if not directory.is_dir():
            raise ValueError("Reconciliation directory does not exist")
        for path in sorted(directory.glob("*.json")):
            content = path.read_bytes()
            part = json.loads(content)
            if not isinstance(part, list):
                raise ValueError("Reconciliation part must be a JSON array: " + path.name)
            part_sources.append({"path": str(path.resolve()), "sha256": hashlib.sha256(content).hexdigest()})
            for evidence in part:
                if not isinstance(evidence, dict):
                    raise ValueError("Reconciliation entry must be an object")
                identity = _requirement_id(evidence.get("requirement_id"))
                if not 1 <= identity <= len(rows):
                    raise ValueError("Reconciliation ID outside original scope: " + str(identity))
                if identity in seen:
                    raise ValueError("Duplicate reconciliation requirement ID: " + str(identity))
                seen.add(identity)
                if not isinstance(evidence.get("source_status"), str) or not evidence["source_status"].strip():
                    raise ValueError("Reconciliation requires an explicit source_status")
                for field in ("source_status", "source_paths", "source_evidence", "missing_capability", "required_access"):
                    if field in evidence:
                        rows[identity - 1][field] = evidence[field]
    report["scope_document_supplied"] = True
    report["scope_document"] = document
    report["original_requirement_rows"] = rows
    report["original_requirement_count"] = len(rows)
    report["source_reviewed_requirement_count"] = len(seen)
    report["source_review_parts"] = part_sources
    report["requirements_reconciled"] = False
    report["limitations"] = [
        "All supplied original requirements are preserved. Source review evidence is separate from live registry reconciliation, deployment and functional verification.",
        "Source-review evidence in optional JSON parts is analyst assessment, not independently executed tests. All production verification fields remain unverified.",
        *[value for value in LIMITATIONS if not value.startswith("The original 161-item")],
    ]
    return report


def _spreadsheet_value(value):
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        value = json.dumps(value, sort_keys=True)
    # Source-derived identifiers must never become spreadsheet formulae.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


TABLES = {"original_requirement_rows": "Requirement source review",
          "edge_functions": "Edge function sources", "action_candidates": "Action source candidates",
          "gateway_handlers": "Gateway handler sources", "function_references": "Function references",
          "missing_worker_sources": "Missing worker sources", "migrations": "Migration sources",
          "source_files": "Source hashes"}


def write_reports(report, output, include_xlsx=False):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    json_path = output / "sandy-api-source-inventory.json"
    json_path.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    written = [json_path]
    for key in TABLES:
        rows = report[key]
        path = output / (key + ".csv")
        fields = list(rows[0]) if rows else ["evidence_status"]
        with path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows({key: _spreadsheet_value(value) for key, value in row.items()}
                             for row in rows)
        written.append(path)
    if include_xlsx:
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError("XLSX requested but openpyxl is not installed; JSON and CSV were written") from exc
        workbook = Workbook()
        metadata = workbook.active
        metadata.title = "Read me and blockers"
        metadata.append(["Evidence boundary", "Value"])
        for key in ("generated_at_utc", "source_head", "source_root", "inventory_scope",
                    "scope_document_supplied", "requirements_reconciled", "production_verified",
                    "original_requirement_count", "source_reviewed_requirement_count", "scope_document",
                    "source_review_parts", "working_tree_status", "counts"):
            metadata.append([key, _spreadsheet_value(report.get(key))])
        for limitation in report["limitations"]:
            metadata.append(["Limitation", limitation])
        for key, title in TABLES.items():
            sheet = workbook.create_sheet(title)
            rows = report[key]
            fields = list(rows[0]) if rows else ["evidence_status"]
            sheet.append(fields)
            for row in rows:
                sheet.append([_spreadsheet_value(row.get(field)) for field in fields])
        for sheet in workbook:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = max(len(str(cell.value or "")) for cell in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 15), 90)
        filename = "sandy-api-reconciliation.xlsx" if report["scope_document_supplied"] else "sandy-api-source-inventory.xlsx"
        path = output / filename
        workbook.save(path)
        written.append(path)
    return written


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--xlsx", action="store_true", help="Also write an Excel source inventory; requires openpyxl")
    parser.add_argument("--scope", type=Path, help="Canonical api-scope-161.md or original Sandy_API_Function_Inventory.xlsx; never modified")
    parser.add_argument("--reconciliation-dir", type=Path, help="JSON source-review arrays keyed by original requirement_id")
    parser.add_argument("--expected-requirements", type=int, default=161)
    args = parser.parse_args(argv)
    report = inventory(args.root)
    if args.reconciliation_dir and not args.scope:
        parser.error("--reconciliation-dir requires --scope")
    if args.scope:
        attach_scope(report, args.scope, args.reconciliation_dir, args.expected_requirements)
    paths = write_reports(report, args.output_dir, args.xlsx)
    print(json.dumps({"files_written": [str(path) for path in paths], "counts": report["counts"],
                      "original_requirement_count": report["original_requirement_count"],
                      "requirements_reconciled": False, "production_verified": False}, sort_keys=True))


if __name__ == "__main__":
    main()
