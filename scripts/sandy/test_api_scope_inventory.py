"""Evidence-boundary tests for the offline API source inventory."""
import csv
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from api_scope_inventory import action_candidates, attach_scope, gateway_handlers, inventory, write_reports


class ApiScopeInventoryTests(unittest.TestCase):
    def test_canonical_markdown_preserves_original_text_and_uses_nearest_domain_without_openpyxl(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "supabase/functions").mkdir(parents=True)
            scope = root / "api-scope.md"
            content = ("# API scope\n\n## Core orchestration (1–1)\n"
                       "1. GET /v1/capabilities — executable operations, scopes and verification.\n"
                       "## Google Cloud organization (2–2)\n"
                       "2. gcp.organizations.list/get\n")
            scope.write_text(content)
            with patch.dict(sys.modules, {"openpyxl": None}):
                report = attach_scope(inventory(root), scope, expected_count=2)
            first, second = report["original_requirement_rows"]
            self.assertEqual(first["domain"], "Core orchestration (1–1)")
            self.assertEqual(first["proposed_api_operation_or_function_family"], "GET /v1/capabilities")
            self.assertEqual(first["required_capability"], "executable operations, scopes and verification.")
            self.assertEqual(first["original_requirement_text"], content.splitlines()[3])
            self.assertEqual(first["source_line"], 4)
            self.assertEqual(second["domain"], "Google Cloud organization (2–2)")
            self.assertEqual(second["required_capability"], "")
            self.assertFalse(first["production_verified"])
            self.assertEqual(report["scope_document"]["format"], "markdown")
            self.assertEqual(report["scope_document"]["sha256"], hashlib.sha256(scope.read_bytes()).hexdigest())

    def test_markdown_rejects_duplicate_skipped_and_wrong_total_requirements(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "supabase/functions").mkdir(parents=True)
            scope = root / "api-scope.md"
            cases = [("## Core\n1. first\n1. repeated\n", "Duplicate Markdown requirement ID"),
                     ("## Core\n1. first\n3. skipped\n", "Expected Markdown requirement ID 2, found 3"),
                     ("## Core\n1. first\n", "Expected 2 original requirements"),
                     ("1. no-domain\n2. still-no-domain\n", "no preceding level-two domain heading")]
            for content, error in cases:
                with self.subTest(error=error):
                    scope.write_text(content)
                    with self.assertRaisesRegex(ValueError, error):
                        attach_scope(inventory(root), scope, expected_count=2)

    def test_action_checks_exclude_comments_and_inequality(self):
        source = '''// if(action === "commented"){}
/* if(body.action == 'also_commented'){} */
if(action === "execute"){}
if(input.action == 'request'){}
if("list" === body.action){}
if(action !== "deny"){}
if(other.action === "unrelated"){}
'''
        self.assertEqual(action_candidates(source), [
            {"action_candidate": "execute", "line": 3},
            {"action_candidate": "request", "line": 4},
            {"action_candidate": "list", "line": 5}])

    def test_gateway_map_is_not_inferred_from_unrelated_objects(self):
        source = '''const unrelated = {"fake.operation": {functionName:"fake",action:"do"}};
const allowedHandlers: Record<string, {functionName:string;action:string}> = {
 "engineering.preflight": {functionName:"sandy-engineering-control",action:"preflight"},
};
'''
        rows = gateway_handlers(source)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["operation"], "engineering.preflight")
        self.assertEqual(rows[0]["handler_action"], "preflight")
        self.assertIn("not_deployed", rows[0]["evidence_status"])
        self.assertEqual(gateway_handlers(source.replace("= {", "= Object.freeze({").replace("\n};", "\n});")), rows)
        self.assertEqual(gateway_handlers(source.replace('action:"preflight"}', 'action:"preflight",\n}'))[0]["handler_action"], "preflight")

    def test_inventory_hashes_sources_and_reports_missing_worker_without_inventing_scope(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "supabase/functions/sandy-test"
            folder.mkdir(parents=True)
            source = b'if(action === "run") fetch(base+"/functions/v1/sandy-absent-worker");'
            (folder / "index.ts").write_bytes(source)
            (folder / "index_test.ts").write_text('if(action === "fixture") fetch(base+"/functions/v1/fake-worker");')
            migration = root / "supabase/migrations/001.sql"
            migration.parent.mkdir()
            migration.write_text("select 1;\n")
            report = inventory(root)
            self.assertEqual(report["counts"]["edge_functions"], 1)
            self.assertEqual(report["counts"]["action_candidates"], 1)
            self.assertEqual(report["counts"]["missing_worker_source_references"], 1)
            self.assertEqual(report["edge_functions"][0]["sha256"], hashlib.sha256(source).hexdigest())
            self.assertEqual(report["missing_worker_sources"][0]["target_function"], "sandy-absent-worker")
            self.assertFalse(report["missing_worker_sources"][0]["target_source_present"])
            self.assertEqual(report["original_requirement_rows"], [])
            self.assertIsNone(report["original_requirement_count"])
            self.assertFalse(report["requirements_reconciled"])
            self.assertFalse(report["production_verified"])
            self.assertEqual(len(report["migrations"]), 1)

    def test_symlink_outside_repository_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp, tempfile.TemporaryDirectory() as external:
            root = Path(temp)
            folder = root / "supabase/functions/test"
            folder.mkdir(parents=True)
            target = Path(external) / "secret.ts"
            target.write_text("not to be read")
            (folder / "index.ts").symlink_to(target)
            with self.assertRaisesRegex(ValueError, "symlink"):
                inventory(root)

    def test_reports_preserve_evidence_and_escape_spreadsheet_formula(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "supabase/functions").mkdir(parents=True)
            report = inventory(root)
            report["edge_functions"] = [{"function_slug": "=DANGEROUS()", "evidence_status": "source_only"}]
            output = root / "reports"
            write_reports(report, output)
            saved = json.loads((output / "sandy-api-source-inventory.json").read_text())
            self.assertFalse(saved["production_verified"])
            with (output / "edge_functions.csv").open(newline="") as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(rows[0]["function_slug"], "'=DANGEROUS()")
            self.assertEqual(rows[0]["evidence_status"], "source_only")

    def test_excel_is_readable_and_carries_missing_scope_blocker(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("Optional XLSX dependency openpyxl is not installed")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "supabase/functions").mkdir(parents=True)
            report = inventory(root)
            report["edge_functions"] = [{"function_slug": "=DANGEROUS()", "evidence_status": "source_only"}]
            output = root / "reports"
            write_reports(report, output, include_xlsx=True)
            workbook = load_workbook(output / "sandy-api-source-inventory.xlsx")
            values = list(workbook["Read me and blockers"].values)
            self.assertIn(("requirements_reconciled", False), values)
            self.assertTrue(any("original 161-item" in str(row) for row in values))
            self.assertEqual(workbook["Edge function sources"]["A2"].value, "'=DANGEROUS()")
            self.assertNotEqual(workbook["Edge function sources"]["A2"].data_type, "f")
            workbook.close()

    def test_original_workbook_text_is_preserved_and_source_claims_cannot_verify_production(self):
        try:
            from openpyxl import Workbook, load_workbook
        except ImportError:
            self.skipTest("Optional scope workbook dependency openpyxl is not installed")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "supabase/functions").mkdir(parents=True)
            scope = root / "scope.xlsx"
            workbook = Workbook()
            workbook.active.title = "API Function Catalogue"
            workbook.active.append(["Domain", "Proposed API operation / function family", "Required capability", "Classification"])
            workbook.active.append(["Original domain", " GET /v1/operations ", "Original capability", "Proposed specification"])
            workbook.save(scope)
            original_bytes = scope.read_bytes()
            parts = root / "parts"
            parts.mkdir()
            evidence = [{"requirement_id": "REQ-001", "source_status": "partial_source",
                         "source_paths": ["supabase/functions/sandy-test/index.ts"],
                         "source_evidence": "Source path only", "production_verified": True,
                         "production_deployment": "verified", "domain": "Untrusted replacement"}]
            (parts / "part.json").write_text(json.dumps(evidence))
            report = attach_scope(inventory(root), scope, parts, expected_count=1)
            row = report["original_requirement_rows"][0]
            self.assertEqual(row["domain"], "Original domain")
            self.assertEqual(row["proposed_api_operation_or_function_family"], " GET /v1/operations ")
            self.assertEqual(row["source_status"], "partial_source")
            self.assertEqual(row["production_deployment"], "unverified")
            self.assertFalse(row["production_verified"])
            self.assertFalse(report["requirements_reconciled"])
            self.assertEqual(report["source_reviewed_requirement_count"], 1)
            self.assertEqual(scope.read_bytes(), original_bytes)
            self.assertEqual(row["scope_document_sha256"], hashlib.sha256(original_bytes).hexdigest())
            output = root / "reports"
            write_reports(report, output, include_xlsx=True)
            generated = load_workbook(output / "sandy-api-reconciliation.xlsx")
            self.assertEqual(generated["Requirement source review"].max_row, 2)
            generated.close()
            (parts / "duplicate.json").write_text(json.dumps(evidence))
            with self.assertRaisesRegex(ValueError, "Duplicate reconciliation"):
                attach_scope(inventory(root), scope, parts, expected_count=1)
            with self.assertRaisesRegex(ValueError, "Expected 161"):
                attach_scope(inventory(root), scope)


if __name__ == "__main__":
    unittest.main()
